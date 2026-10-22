from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Administrator\Desktop\work\停车场管理数据库项目")
WORK = Path(r"C:\Users\Administrator\Desktop\work\_report_build")
SOURCE = WORK / "template.docx"
OUTPUT = ROOT / "停车场管理项目报告.docx"
ENTITY_IMAGE = WORK / "entity-attributes.png"
ER_IMAGE = WORK / "overall-er.png"


# Database source identifiers use English snake_case; Chinese labels remain in
# the report narrative and diagrams for readability.
SOURCE_IDENTIFIERS = {
    "管理员": "administrator", "车主": "vehicle_owner", "车辆": "vehicle",
    "车位": "parking_space", "停车记录": "parking_record", "收费规则": "charge_rule",
    "月租合同": "rental_contract", "缴费记录": "payment_record",
    "管理员ID": "administrator_id", "工号": "employee_no", "姓名": "name",
    "联系电话": "phone", "登录账号": "login_account", "角色": "role",
    "账号状态": "account_status", "创建时间": "created_at", "车主ID": "owner_id",
    "身份证号": "id_number", "联系地址": "address", "车主状态": "owner_status",
    "车辆ID": "vehicle_id", "车牌号": "license_plate", "车辆类型": "vehicle_type",
    "车辆颜色": "vehicle_color", "车辆状态": "vehicle_status", "登记时间": "registered_at",
    "车位ID": "parking_space_id", "车位编号": "space_code", "区域": "area",
    "车位类型": "space_type", "车位状态": "space_status", "位置说明": "location_note",
    "停车记录ID": "parking_record_id", "停车类型": "parking_type",
    "收费规则ID": "charge_rule_id", "月租合同ID": "rental_contract_id",
    "入场时间": "entry_time", "出场时间": "exit_time", "停车时长分钟": "duration_minutes",
    "应收金额": "amount_due", "记录状态": "record_status", "办理管理员ID": "handled_by_administrator_id",
    "备注": "remark", "规则名称": "rule_name", "每小时单价": "hourly_rate",
    "最低收费": "minimum_charge", "每日封顶金额": "daily_cap", "生效日期": "effective_date",
    "失效日期": "expiry_date", "规则状态": "rule_status", "创建管理员ID": "created_by_administrator_id",
    "合同编号": "contract_no", "开始日期": "start_date", "结束日期": "end_date",
    "月租金额": "monthly_rent", "押金": "deposit", "合同状态": "contract_status",
    "签订时间": "signed_at", "支付流水号": "payment_no", "缴费类型": "payment_type",
    "实收金额": "amount_paid", "支付方式": "payment_method", "支付状态": "payment_status",
    "支付时间": "paid_at", "车主姓名": "owner_name", "支付日期": "payment_date",
    "缴费笔数": "payment_count", "收入金额": "revenue_amount", "已支付金额": "paid_amount",
    "已支付次数": "paid_count", "缴费记录ID": "payment_record_id",
    "办理管理员": "administrator_name",
}


def source_identifier(value):
    return SOURCE_IDENTIFIERS.get(value, value)


def translate_sql(sql):
    for chinese, english in SOURCE_IDENTIFIERS.items():
        sql = sql.replace(f"`{chinese}`", f"`{english}`")
    return sql


def set_run(run, size=12, bold=None, font="宋体", color="000000", underline=None):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline


def replace_cell(cell, value):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    run = p.add_run(value + " " * max(0, 42 - len(value)))
    set_run(run, size=11, bold=True if value == "待补" else None, underline=True)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def body_paragraph(doc, text, indent=True, size=12, bold_lead=None, keep_with_next=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run(r, size=size, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run(r, size=size)
    else:
        r = p.add_run(text)
        set_run(r, size=size)
    return p


def heading(doc, text, level, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def caption(doc, text, page_break=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(text)
    set_run(r, size=10.5)
    return p


def add_figure(doc, path, width=6.15, page_break=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    if page_break:
        p.paragraph_format.page_break_before = True
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def add_sql(doc, sql):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = Pt(11.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    r = p.add_run(translate_sql(sql.strip()))
    set_run(r, size=8.5, font="Consolas")
    return p


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_physical_table(doc, number, name, rows):
    p = doc.add_paragraph(f"{number}. {name}表", style="Normal")
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = number in {1, 3, 5, 7, 8}
    for r in p.runs:
        set_run(r, size=12)
    caption(doc, f"表 2-{number}  {name}表字段设计")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Normal Table"
    headers = ["字段", "类型", "长度", "约束", "主外键", "备注"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            for run in para.runs:
                set_run(run, size=8.5, bold=True)
    repeat_table_header(table.rows[0])
    widths = [0.84, 1.07, 0.50, 1.25, 0.62, 1.88]
    for row in rows:
        row = list(row)
        row[0] = source_identifier(row[0])
        if row[4].startswith("FK→"):
            row[4] = "FK→" + source_identifier(row[4][3:])
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cells[i].paragraphs:
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = Pt(10.5)
                if i in (0, 1, 2, 4):
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_run(run, size=8.8)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
    doc.add_paragraph()
    return table


TABLES = [
    ("管理员", [
        ("administrator_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "管理员内部标识"),
        ("employee_no", "VARCHAR", "20", "NOT NULL, UNIQUE", "—", "员工业务编号"),
        ("name", "VARCHAR", "50", "NOT NULL", "—", "管理员姓名"),
        ("phone", "VARCHAR", "20", "NOT NULL, UNIQUE", "—", "联系方式"),
        ("login_account", "VARCHAR", "30", "NOT NULL, UNIQUE", "—", "系统登录账号"),
        ("role", "VARCHAR", "20", "NOT NULL, CHECK", "—", "收费员/主管/系统管理员"),
        ("account_status", "VARCHAR", "10", "NOT NULL, DEFAULT '启用', CHECK", "—", "启用/停用"),
        ("created_at", "DATETIME", "—", "NOT NULL, DEFAULT CURRENT_TIMESTAMP", "—", "账号创建时间"),
    ]),
    ("车主", [
        ("owner_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "车主内部标识"),
        ("name", "VARCHAR", "50", "NOT NULL", "—", "车主姓名"),
        ("id_number", "CHAR", "18", "NOT NULL, UNIQUE", "—", "实名识别"),
        ("phone", "VARCHAR", "20", "NOT NULL, UNIQUE", "—", "联系方式"),
        ("address", "VARCHAR", "200", "NULL", "—", "可选地址"),
        ("owner_status", "VARCHAR", "10", "NOT NULL, DEFAULT '正常', CHECK", "—", "正常/停用"),
        ("created_at", "DATETIME", "—", "NOT NULL, DEFAULT CURRENT_TIMESTAMP", "—", "登记时间"),
    ]),
    ("车位", [
        ("parking_space_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "车位内部标识"),
        ("space_code", "VARCHAR", "20", "NOT NULL, UNIQUE", "—", "现场唯一标识"),
        ("area", "VARCHAR", "30", "NOT NULL", "—", "车位所属区域"),
        ("space_type", "VARCHAR", "20", "NOT NULL, CHECK", "—", "标准/无障碍/充电"),
        ("space_status", "VARCHAR", "10", "NOT NULL, DEFAULT '可用', CHECK", "—", "可用/占用/停用"),
        ("location_note", "VARCHAR", "100", "NULL", "—", "位置补充说明"),
    ]),
    ("车辆", [
        ("vehicle_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "车辆内部标识"),
        ("license_plate", "VARCHAR", "12", "NOT NULL, UNIQUE", "—", "本停车场内唯一"),
        ("vehicle_type", "VARCHAR", "20", "NOT NULL, CHECK", "—", "小型车/大型车/新能源"),
        ("vehicle_color", "VARCHAR", "20", "NULL", "—", "外观颜色"),
        ("owner_id", "BIGINT UNSIGNED", "—", "NOT NULL, INDEX", "FK→vehicle_owner", "所属车主"),
        ("vehicle_status", "VARCHAR", "10", "NOT NULL, DEFAULT '正常', CHECK", "—", "正常/停用"),
        ("registered_at", "DATETIME", "—", "NOT NULL, DEFAULT CURRENT_TIMESTAMP", "—", "车辆登记时间"),
    ]),
    ("收费规则", [
        ("charge_rule_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "规则标识"),
        ("rule_name", "VARCHAR", "50", "NOT NULL, UNIQUE", "—", "收费规则名称"),
        ("hourly_rate", "DECIMAL", "10,2", "NOT NULL, CHECK > 0", "—", "临停小时单价"),
        ("minimum_charge", "DECIMAL", "10,2", "NOT NULL, DEFAULT 0.00, CHECK >= 0", "—", "最低应收金额"),
        ("daily_cap", "DECIMAL", "10,2", "NULL, CHECK为空或>0", "—", "可选每日封顶"),
        ("effective_date", "DATE", "—", "NOT NULL", "—", "规则开始日期"),
        ("expiry_date", "DATE", "—", "NULL, CHECK为空或不早于生效日", "—", "为空表示仍有效"),
        ("rule_status", "VARCHAR", "10", "NOT NULL, DEFAULT '启用', CHECK", "—", "启用/停用"),
        ("created_by_administrator_id", "BIGINT UNSIGNED", "—", "NOT NULL, INDEX", "FK→administrator", "规则创建人员"),
    ]),
    ("月租合同", [
        ("rental_contract_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "合同标识"),
        ("contract_no", "VARCHAR", "30", "NOT NULL, UNIQUE", "—", "合同业务编号"),
        ("vehicle_id", "BIGINT UNSIGNED", "—", "NOT NULL, INDEX", "FK→vehicle", "合同所属车辆"),
        ("start_date", "DATE", "—", "NOT NULL", "—", "合同开始日"),
        ("end_date", "DATE", "—", "NOT NULL, CHECK >=开始日期", "—", "合同结束日"),
        ("monthly_rent", "DECIMAL", "10,2", "NOT NULL, CHECK > 0", "—", "每月应缴金额"),
        ("deposit", "DECIMAL", "10,2", "NOT NULL, DEFAULT 0.00, CHECK >= 0", "—", "可选押金"),
        ("contract_status", "VARCHAR", "10", "NOT NULL, DEFAULT '生效', CHECK", "—", "生效/到期/终止"),
        ("signed_at", "DATETIME", "—", "NOT NULL, DEFAULT CURRENT_TIMESTAMP", "—", "业务时间"),
        ("handled_by_administrator_id", "BIGINT UNSIGNED", "—", "NOT NULL", "FK→administrator", "办理合同人员"),
        ("remark", "VARCHAR", "200", "NULL", "—", "合同说明"),
    ]),
    ("停车记录", [
        ("parking_record_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "一次进出场标识"),
        ("vehicle_id", "BIGINT UNSIGNED", "—", "NOT NULL, INDEX", "FK→vehicle", "停放车辆"),
        ("parking_space_id", "BIGINT UNSIGNED", "—", "NOT NULL, INDEX", "FK→parking_space", "占用车位"),
        ("parking_type", "VARCHAR", "10", "NOT NULL, CHECK", "—", "临时停车/月租停车"),
        ("charge_rule_id", "BIGINT UNSIGNED", "—", "NULL, 按停车类型互斥", "FK→charge_rule", "临停必填、月租为空"),
        ("rental_contract_id", "BIGINT UNSIGNED", "—", "NULL, 按停车类型互斥", "FK→rental_contract", "月租必填、临停为空"),
        ("entry_time", "DATETIME", "—", "NOT NULL", "—", "车辆入场时间"),
        ("exit_time", "DATETIME", "—", "NULL, CHECK为空或不早于入场", "—", "在场时为空"),
        ("duration_minutes", "INT UNSIGNED", "—", "NULL, CHECK为空或>=0", "—", "出场时计算"),
        ("amount_due", "DECIMAL", "10,2", "NOT NULL, DEFAULT 0.00, CHECK >= 0", "—", "月租停车固定为0"),
        ("record_status", "VARCHAR", "10", "NOT NULL, DEFAULT '在场', CHECK", "—", "在场/已出场/异常"),
        ("handled_by_administrator_id", "BIGINT UNSIGNED", "—", "NOT NULL", "FK→administrator", "登记业务人员"),
        ("remark", "VARCHAR", "200", "NULL", "—", "异常或人工说明"),
    ]),
    ("缴费记录", [
        ("payment_record_id", "BIGINT UNSIGNED", "—", "NOT NULL, AUTO_INCREMENT", "PK", "缴费标识"),
        ("payment_no", "VARCHAR", "40", "NOT NULL, UNIQUE", "—", "外部/内部流水号"),
        ("payment_type", "VARCHAR", "10", "NOT NULL, CHECK并与归属一致", "—", "临停停车费/月租费"),
        ("parking_record_id", "BIGINT UNSIGNED", "—", "NULL, 与合同ID二选一", "FK→parking_record", "临停缴费归属"),
        ("rental_contract_id", "BIGINT UNSIGNED", "—", "NULL, 与停车记录ID二选一", "FK→rental_contract", "月租缴费归属"),
        ("amount_due", "DECIMAL", "10,2", "NOT NULL, CHECK >= 0", "—", "业务应收金额"),
        ("amount_paid", "DECIMAL", "10,2", "NOT NULL, CHECK >= 0；状态联动", "—", "待支付为0；已支付/已退款>0"),
        ("payment_method", "VARCHAR", "10", "NOT NULL, CHECK", "—", "现金/微信/支付宝/银行卡"),
        ("payment_status", "VARCHAR", "10", "NOT NULL, DEFAULT '已支付', CHECK", "—", "待支付/已支付/已退款"),
        ("paid_at", "DATETIME", "—", "NULL, DEFAULT NULL, INDEX；状态联动", "—", "待支付为空；其余为实际支付时间"),
        ("handled_by_administrator_id", "BIGINT UNSIGNED", "—", "NOT NULL", "FK→administrator", "登记收款人员"),
        ("remark", "VARCHAR", "200", "NULL", "—", "退款或人工说明"),
    ]),
]


QUERIES = [
    ("单表查询1：当前在场的临时车辆", "筛选尚未出场且类型为临时停车的记录，供管理员核对现场放行车辆。", """
SELECT `parking_record_id`, `vehicle_id`, `parking_space_id`, `entry_time`, `record_status`
FROM `parking_record`
WHERE `parking_type` = '临时停车'
  AND `record_status` = '在场'
  AND `exit_time` IS NULL
ORDER BY `entry_time`;
"""),
    ("单表查询2：未来30天内到期的有效月租合同", "使用动态日期范围生成续费提醒清单。", """
SELECT `rental_contract_id`, `contract_no`, `vehicle_id`, `start_date`, `end_date`, `monthly_rent`
FROM `rental_contract`
WHERE `contract_status` = '生效'
  AND `end_date` BETWEEN CURRENT_DATE()
      AND DATE_ADD(CURRENT_DATE(), INTERVAL 30 DAY)
ORDER BY `end_date`;
"""),
    ("单表查询3：指定日期已支付流水", "用半开区间覆盖指定日期全天，筛选已支付收款用于日结核对。", """
SELECT `payment_no`, `payment_type`, `amount_due`, `amount_paid`, `payment_method`, `paid_at`
FROM `payment_record`
WHERE `payment_status` = '已支付'
  AND `paid_at` >= '2026-10-19 00:00:00'
  AND `paid_at` <  '2026-10-20 00:00:00'
ORDER BY `paid_at`;
"""),
    ("多表查询1：在场车辆、车主和车位", "连接停车记录、车辆、车主和车位，显示实时占用情况。", """
SELECT v.`license_plate`, o.`name` AS `owner_name`,
       ps.`space_code`, ps.`area`, pr.`entry_time`, pr.`parking_type`
FROM `parking_record` pr
JOIN `vehicle` v ON v.`vehicle_id` = pr.`vehicle_id`
JOIN `vehicle_owner` o ON o.`owner_id` = v.`owner_id`
JOIN `parking_space` ps ON ps.`parking_space_id` = pr.`parking_space_id`
WHERE pr.`record_status` = '在场'
  AND pr.`exit_time` IS NULL
ORDER BY ps.`area`, ps.`space_code`;
"""),
    ("多表查询2：按月租合同统计已支付金额", "左连接保留尚无付款的合同，条件聚合仅累计已支付流水。", """
SELECT mc.`contract_no`, v.`license_plate`, o.`name` AS `owner_name`,
       mc.`monthly_rent`,
       COALESCE(SUM(CASE WHEN p.`payment_status` = '已支付'
                         THEN p.`amount_paid` ELSE 0 END), 0) AS `paid_amount`,
       COUNT(CASE WHEN p.`payment_status` = '已支付'
                  THEN p.`payment_record_id` END) AS `paid_count`
FROM `rental_contract` mc
JOIN `vehicle` v ON v.`vehicle_id` = mc.`vehicle_id`
JOIN `vehicle_owner` o ON o.`owner_id` = v.`owner_id`
LEFT JOIN `payment_record` p ON p.`rental_contract_id` = mc.`rental_contract_id`
GROUP BY mc.`rental_contract_id`, mc.`contract_no`, v.`license_plate`, o.`name`, mc.`monthly_rent`
ORDER BY mc.`end_date`;
"""),
    ("多表查询3：按日期、缴费类型和管理员统计收入", "连接管理员并按日期、类型、经办人分组，统计笔数和实收金额。", """
SELECT DATE(p.`paid_at`) AS `payment_date`,
       p.`payment_type`, a.`name` AS `administrator_name`,
       COUNT(*) AS `payment_count`, SUM(p.`amount_paid`) AS `revenue_amount`
FROM `payment_record` p
JOIN `administrator` a ON a.`administrator_id` = p.`handled_by_administrator_id`
WHERE p.`payment_status` = '已支付'
GROUP BY DATE(p.`paid_at`), p.`payment_type`, a.`administrator_id`, a.`name`
ORDER BY `payment_date` DESC, p.`payment_type`, a.`administrator_id`;
"""),
]


def clear_after_start(doc):
    body = doc._body._element
    start = None
    for child in list(body):
        if child.tag == qn("w:p"):
            text = "".join(child.itertext())
            if "本文档蓝色部分文字内容" in text:
                start = child
                break
    if start is None:
        raise RuntimeError("Template body start marker was not found")
    remove = False
    for child in list(body):
        if child is start:
            remove = True
        if child.tag == qn("w:sectPr"):
            continue
        if remove:
            body.remove(child)


def main():
    doc = Document(SOURCE)

    cover = doc.tables[0]
    replace_cell(cover.cell(1, 1), "停车场管理系统")
    replace_cell(cover.cell(4, 1), "待补")
    replace_cell(cover.cell(5, 1), "待补")

    clear_after_start(doc)

    heading(doc, "一、项目概述", 3)
    body_paragraph(doc, "本项目面向单个停车场，围绕车主与车辆登记、车位管理、车辆进出场、临时停车计费、月租合同、缴费记录及运营统计开展数据库设计。项目以 MySQL 8.0 为实现环境，完成需求分析、概念模型、关系模式和物理表结构设计，并提供可执行的建库建表及样例数据脚本和六个核心查询。")

    heading(doc, "二、数据库设计全流程", 3)
    heading(doc, "2.1 需求分析", 4)
    body_paragraph(doc, "系统服务于车主、停车场管理员和数据库应用服务。设计范围限于一个停车场，不包括车位预约、跨场调度、优惠券、发票和多停车场管理。数据模型围绕车辆进出场及费用结算组织，历史业务记录保留，避免通过删除主数据破坏追溯关系。")
    body_paragraph(doc, "（1）车主需求：登记姓名、身份证号、联系电话和联系地址，维护名下车辆，并查询车辆停车与缴费情况；符合停车场规则的车辆可办理月租合同。一个车主可以登记多辆车，每辆车在本场仅归属一个车主。")
    body_paragraph(doc, "（2）管理员需求：维护车主、车辆、车位、收费规则及管理员账号；登记车辆入场和出场，办理月租合同，登记临停或月租缴费及退款；查询在场车辆、占用车位、车辆停车历史、合同到期情况、分类收入和经办业务量。每条停车记录、月租合同、缴费记录及收费规则都保留对应的经办或创建管理员。")
    body_paragraph(doc, "（3）系统处理需求：校验车牌、车位、合同有效期、金额和记录状态；临停按规则以不足一小时进一小时计算，月租合同有效期内的停车记录应收金额为零；车辆出场后计算停车时长和应收金额，并释放车位。系统应支持按车辆、日期、收费类型、合同和管理员开展查询统计。")
    body_paragraph(doc, "（4）主要业务流程：先登记车主及车辆、车位和收费规则；车辆入场时校验同一车辆无未出场记录且目标车位可用，建立停车记录并更新车位状态；月租车辆关联当前有效合同，临时车辆关联当时适用的收费规则；出场时检查出场时间不早于入场时间，计算临停金额或将月租停车金额记为零；收费时记录支付方式、金额、时间、状态及唯一业务归属；退款通过更新支付状态保留原始流水。")
    body_paragraph(doc, "（5）核心业务约束：车牌号、车位编号、合同编号、支付流水号及管理员工号等业务标识唯一；停车记录的收费规则与月租合同按停车类型互斥；每笔缴费必须且只能关联一条停车记录或一份月租合同，缴费类型须与归属相符；同一车辆或车位不能同时存在多条未出场记录，同一车辆的有效月租合同日期不能重叠。最后三项涉及多行或日期区间，需在事务、触发器或应用服务中校验。")
    body_paragraph(doc, "（6）主要数据范围：系统设置车主、车辆、车位、停车记录、收费规则、月租合同、缴费记录和管理员八个实体。姓名、证件、联系方式、车牌、区域、入出场时间、收费参数、合同期限、支付流水和处理人员等数据均由对应实体存储，避免在业务表中重复保存可由外键取得的主数据。")

    heading(doc, "2.2 概念结构设计（ER图设计）", 4)
    body_paragraph(doc, "概念模型包含八个实体。实体属性图展示每个实体的主键及主要业务属性；整体图使用 Chen 记法表达实体、联系菱形和联系基数。完整字段以本报告物理表设计为准。")
    heading(doc, "（1）实体属性图", 5)
    body_paragraph(doc, "图中椭圆表示属性，标注 PK 的字段为主键，标注 FK 的字段为外键；图示聚焦核心属性，实体的完整字段及约束见表 2-1 至表 2-8。")
    caption(doc, "图 2-1  停车场管理系统实体属性图")
    add_figure(doc, ENTITY_IMAGE, width=6.15)
    heading(doc, "（2）实体及实体间联系", 5)
    body_paragraph(doc, "车主与车辆为 1:N 联系：一个车主可以登记多辆车，每辆车仅属于一个车主。车辆与停车记录为 1:N 联系：一辆车可产生多次停车记录，每条记录只对应一辆车。车位与停车记录为 1:N 联系：车位在不同时段可承载多次停车记录，每条记录占用一个车位。")
    body_paragraph(doc, "车辆与月租合同为 1:N 联系：一辆车可在不同时间签订多份合同，每份合同对应一辆车；同一车辆有效合同不得重叠。收费规则与停车记录为 1:N 可选联系：临停记录必须关联一条收费规则，月租记录不关联收费规则。整体图未单独绘制月租合同与停车记录的直接联系，月租缴费通过月租合同与缴费记录之间的“月租缴费”联系表达，临停缴费通过停车记录与缴费记录之间的“临停缴费”联系表达。")
    body_paragraph(doc, "停车记录与缴费记录、月租合同与缴费记录均为 1:N 联系。缴费记录归属两种联系之一，不能同时关联停车记录和月租合同，也不能两者皆空。管理员分别与停车记录、月租合同、缴费记录和收费规则构成 1:N 办理或维护联系。")
    heading(doc, "（3）整体 Chen E-R 图", 5)
    body_paragraph(doc, "整体图采用矩形表示实体、菱形表示联系，连线旁标出 1 与 N 基数。实体属性见图 2-1 和各表字段设计；停车记录收费规则/合同以及缴费归属的互斥条件由数据库 CHECK 约束落实。")
    caption(doc, "图 2-2  停车场管理系统整体 Chen E-R 图")
    add_figure(doc, ER_IMAGE, width=6.15)

    heading(doc, "2.3 逻辑结构设计", 4)
    body_paragraph(doc, "依据实体间联系将概念模型转换为八个关系模式。各模式首项为主键，FK 标记外键及其参照关系。所有实体均采用单列主键；1:N 联系的外键保存在 N 端。")
    relations = [
        "vehicle_owner（owner_id，name，id_number，phone，address，owner_status，created_at）",
        "vehicle（vehicle_id，license_plate，vehicle_type，vehicle_color，owner_id，vehicle_status，registered_at）；owner_id FK→vehicle_owner（owner_id）",
        "parking_space（parking_space_id，space_code，area，space_type，space_status，location_note）",
        "parking_record（parking_record_id，vehicle_id，parking_space_id，parking_type，charge_rule_id，rental_contract_id，entry_time，exit_time，duration_minutes，amount_due，record_status，handled_by_administrator_id，remark）；外键分别参照vehicle、parking_space、charge_rule、rental_contract和administrator",
        "charge_rule（charge_rule_id，rule_name，hourly_rate，minimum_charge，daily_cap，effective_date，expiry_date，rule_status，created_by_administrator_id）；created_by_administrator_id FK→administrator（administrator_id）",
        "rental_contract（rental_contract_id，contract_no，vehicle_id，start_date，end_date，monthly_rent，deposit，contract_status，signed_at，handled_by_administrator_id，remark）；vehicle_id FK→vehicle（vehicle_id），handled_by_administrator_id FK→administrator（administrator_id）",
        "payment_record（payment_record_id，payment_no，payment_type，parking_record_id，rental_contract_id，amount_due，amount_paid，payment_method，payment_status，paid_at，handled_by_administrator_id，remark）；parking_record_id、rental_contract_id 二选一，handled_by_administrator_id FK→administrator（administrator_id）",
        "administrator（administrator_id，employee_no，name，phone，login_account，role，account_status，created_at）",
    ]
    for n, relation in enumerate(relations, 1):
        body_paragraph(doc, f"{n}. {relation}", indent=False, size=11.5)
    heading(doc, "规范化分析", 5)
    body_paragraph(doc, "八个关系的字段均为不可再分的单值数据，满足第一范式。各表非主属性完全依赖本表主键，不存在对复合主键的部分依赖，满足第二范式。车主、车辆、合同、收费规则和管理员资料分别存储，业务记录仅保留标识及必要的业务快照，非主属性不经由其他非主属性传递依赖，满足第三范式。停车记录中的应收金额保留为结算快照，防止收费规则调整后历史账单被重新解释。")

    heading(doc, "2.4 物理结构设计", 4)
    body_paragraph(doc, "项目使用 MySQL 8.0，数据库字符集为 utf8mb4、存储引擎为 InnoDB。八张表的主键均为自增 BIGINT UNSIGNED；金额使用 DECIMAL(10,2) 避免浮点误差；日期和业务时间分别使用 DATE 与 DATETIME。状态和类型用 VARCHAR 并以 CHECK 限定取值。唯一键、外键、普通索引及检查约束与项目 SQL 脚本一致。")
    body_paragraph(doc, "外键采用 ON DELETE RESTRICT，保护被引用的历史业务数据；多数外键更新规则为 ON UPDATE CASCADE。停车记录中的收费规则ID/月租合同ID、缴费记录中的停车记录ID/月租合同ID参与互斥 CHECK，按 MySQL 8.0 对相关外键更新动作的限制使用 ON UPDATE RESTRICT。")
    body_paragraph(doc, "以下表格列出字段类型、长度、空值/默认值/检查等约束、主外键及业务说明。字段名和约束以《停车场管理数据库.sql》建表语句为准。")
    for idx, (name, rows) in enumerate(TABLES, 1):
        add_physical_table(doc, idx, name, rows)
    heading(doc, "索引与跨记录一致性", 5)
    body_paragraph(doc, "脚本为车牌、身份证号、联系电话、车位编号、规则名、合同编号、支付流水号、工号和登录账号建立唯一约束；为 vehicle 的 owner_id、parking_record 的（vehicle_id，entry_time）与（parking_space_id，entry_time）、record_status、rental_contract 的 vehicle_id 与日期、paid_at/类型/状态及缴费归属外键建立查询索引。索引支持常用查询和候选冲突检查。")
    body_paragraph(doc, "同一车辆或车位最多一条在场记录，以及同一车辆有效月租合同不重叠，属于跨记录规则，不能由单行 CHECK 完成。业务写入时应在事务中锁定车辆/车位或合同范围，执行冲突检查；入场时同步占用车位，出场时补全记录并释放车位；退款只更新支付状态，不删除流水。")

    heading(doc, "三、核心查询操作实现", 3)
    body_paragraph(doc, "项目设计了三个单表查询和三个多表查询，涵盖条件筛选、日期区间、内连接、左连接、条件聚合、分组统计和排序。SQL 使用与建表脚本一致的英文 snake_case 表名及字段名，业务枚举值保留中文。")
    for title, purpose, sql in QUERIES:
        heading(doc, title, 4)
        body_paragraph(doc, "查询用途：" + purpose, keep_with_next=True)
        add_sql(doc, sql)

    heading(doc, "四、项目能力达成与问题解决", 3)
    body_paragraph(doc, "本项目将业务角色、停车流程、收费规则与数据对象逐项对应，形成八实体模型、实体属性图、整体 Chen E-R 图和关系模式。由概念结构向物理结构转换时，按主外键依赖顺序组织表结构，使实体、关系和字段在模型与 DDL 间保持一致。")
    body_paragraph(doc, "金额字段使用定点小数，时间字段区分业务日期与精确时刻；通过 NOT NULL、UNIQUE、CHECK、外键和索引表达可由数据库直接保证的规则。对于车辆/车位占用以及合同区间重叠等跨记录约束，明确采用事务锁定和冲突校验，使完整性方案与 MySQL 单行约束能力相匹配。")
    body_paragraph(doc, "六个查询覆盖单表筛选、动态日期、指定日期半开区间、多表连接、LEFT JOIN 保留无缴费合同、条件求和与按管理员分组汇总。已支付状态参与收入统计，退款流水不纳入收入；关联多个可能为一对多的数据时先确认分组粒度，避免金额重复累计。")
    body_paragraph(doc, "项目材料包括需求与模型说明、draw.io 图、MySQL 建库建表和样例数据脚本，以及字段设计和查询说明。下一步可在 MySQL 8.0 环境执行脚本，运行六条查询并按课程验收要求留存结果截图，同时验证唯一、检查和参照完整性约束。")

    heading(doc, "五、项目总结", 3)
    body_paragraph(doc, "通过停车场管理数据库项目，完成了从业务范围界定、需求分析、概念结构设计、逻辑结构转换到 MySQL 物理结构设计的全过程。八个实体的属性、联系、关系模式、字段约束与查询均围绕车辆进出场、临停计费、月租及缴费业务展开。")
    body_paragraph(doc, "本项目的关键设计是保留可追溯的停车和缴费历史，并将互斥归属、有效合同、金额和记录状态等业务规则明确落实到数据库约束或事务校验中。查询部分进一步说明如何从规范化数据结构中取得实时占用、合同收款和分类收入信息，为停车场日常管理提供基础数据支持。")

    doc.core_properties.title = "停车场管理系统数据库课程项目报告"
    doc.core_properties.subject = "数据库原理与应用综合项目设计"
    doc.core_properties.author = "待补"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

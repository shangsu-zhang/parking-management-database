from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


root = Path(__file__).resolve().parents[2]
path = root / "report" / "停车场管理项目报告.docx"
doc = Document(path)
cover = doc.tables[0]

for row_index in (4, 5, 6):
    row = cover.rows[row_index]
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(12)

in_summary = False
for paragraph in doc.paragraphs:
    if paragraph.text in ("四、项目能力达成与问题解决", "五、项目总结"):
        in_summary = True
        if paragraph.text == "五、项目总结":
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(6)
        continue
    if in_summary and paragraph.text.strip():
        paragraph.paragraph_format.line_spacing = Pt(18)

doc.save(path)
print(path)
